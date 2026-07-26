"""첨부 파싱 — docx/xlsx 실물 왕복, txt 인코딩 폴백, 예산 클리핑."""

import io

import pytest
from docx import Document
from openpyxl import Workbook

from app.interview.parsing import (
    ALLOWED_EXTENSIONS,
    ParseError,
    clip_to_budget,
    parse_attachment,
)


def _docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_docx_extracts_paragraphs() -> None:
    text = parse_attachment("sop.docx", _docx_bytes(["구매 요청 절차", "1. 요청서 작성"]))
    assert "구매 요청 절차" in text
    assert "요청서 작성" in text


def test_parse_xlsx_extracts_cells_tab_separated() -> None:
    text = parse_attachment("list.xlsx", _xlsx_bytes([["단계", "담당"], ["접수", "구매팀"]]))
    assert "단계\t담당" in text
    assert "접수\t구매팀" in text


def test_parse_txt_utf8_and_cp949() -> None:
    assert parse_attachment("a.txt", "한글 메모".encode("utf-8")) == "한글 메모"
    assert parse_attachment("b.txt", "한글 메모".encode("cp949")) == "한글 메모"


def test_parse_unknown_extension_raises() -> None:
    with pytest.raises(ParseError):
        parse_attachment("evil.exe", b"MZ")


def test_parse_corrupt_docx_raises_parse_error() -> None:
    with pytest.raises(ParseError):
        parse_attachment("broken.docx", b"not a zip")


def test_allowed_extensions() -> None:
    assert ALLOWED_EXTENSIONS == {".pdf", ".docx", ".xlsx", ".txt", ".md"}


def test_clip_to_budget_headers_and_even_cut() -> None:
    sections = [("a.txt", "가" * 100), ("b.txt", "나" * 100)]
    merged = clip_to_budget(sections, budget=120)
    assert "[a.txt]" in merged and "[b.txt]" in merged
    # 각 섹션 본문이 예산의 절반 수준으로 잘림
    assert len(merged) <= 120 + len("[a.txt]\n\n") + len("[b.txt]\n\n") + 4


def test_clip_to_budget_no_cut_when_under() -> None:
    merged = clip_to_budget([("a.txt", "짧다")], budget=1000)
    assert "짧다" in merged
